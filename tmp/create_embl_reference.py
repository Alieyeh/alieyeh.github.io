from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


ROOT = Path("cvs/recommendation_letters")
SOURCE = ROOT / "Fatemeh_Torabi_General_Academic_Reference_Draft.docx"
OUTPUT = ROOT / "Fatemeh_Torabi_EMBL_EBI_Software_Engineer_Data_Discovery_JR3882_Reference_Draft.docx"

doc = Document(SOURCE)

body = [
    "[Date]",
    "Re: Recommendation for Alieyeh Sarabandi Moghaddam - Software Engineer, Data Discovery (JR3882), EMBL-EBI",
    "Dear Members of the Hiring Committee,",
    (
        "I am pleased to recommend Ms Alieyeh Sarabandi Moghaddam for the Software Engineer - "
        "Data Discovery position at EMBL-EBI. I have worked with Alieyeh as a research and writing "
        "supervisor within Dementias Platform UK (DPUK), and have come to know her as an intellectually "
        "curious, conscientious and technically strong colleague whose work sits naturally at the interface "
        "of biomedical data, software engineering and research infrastructure."
    ),
    (
        "The focus of this role - integrating biological data and metadata, supporting FAIR discovery, "
        "and developing reliable software for a broad research community - is very well aligned with "
        "Alieyeh's strengths. In her current work she contributes to systems and workflows that make genomic, "
        "multi-omic and cohort resources more findable, interpretable and usable. She thinks carefully about "
        "identifiers, provenance, metadata completeness, terminology, data quality and the practical needs of "
        "researchers who must assess whether a dataset is suitable before formal access or analysis."
    ),
    (
        "Our most substantial collaboration led to Choosing the Right Genomic Dataset: A Five-Pillar "
        "Framework for Researchers, published by Real World Data Science with Alieyeh as first author. "
        "She took a broad and technically complex topic and developed it into a practical framework "
        "spanning dataset discovery, governance, assay and technology choice, cohort design, quality "
        "control, harmonisation and research readiness. This work demonstrated her ability to synthesise "
        "evidence, identify the decisions that matter to research validity and communicate technical ideas "
        "clearly without losing precision."
    ),
    (
        "Alieyeh's broader DPUK work gives her directly relevant experience for a data discovery platform. "
        "She has worked on reproducible genomic and multi-omic workflows, metadata and provenance, secure data "
        "processing, dataset validation and researcher-facing discovery and feasibility tooling. She has also "
        "designed and supervised work on omics metadata infrastructure, including domain modelling, standards-aware "
        "metadata, validation and search-oriented structures. Her BSc in Computer Engineering and MSc in Health "
        "Data Science allow her to move confidently between software implementation, data engineering, statistical "
        "analysis and biomedical interpretation."
    ),
    (
        "I have been particularly impressed by the way Alieyeh approaches feedback and technical ambiguity. "
        "She listens closely, tests suggestions against the evidence, and revises with purpose rather than "
        "making superficial changes. She is also willing to question assumptions and defend a position when the "
        "scientific reasoning supports it. This balance of receptiveness and independence is important for a "
        "software engineer working in an established, collaborative codebase where design decisions must serve "
        "long-term maintainability and user needs."
    ),
    (
        "I recommend Alieyeh with confidence for the Software Engineer - Data Discovery role. She would "
        "bring strong analytical ability, research maturity, biomedical metadata insight, clear communication "
        "and a collaborative engineering mindset to EMBL-EBI. I expect her to make a thoughtful and valuable "
        "contribution to EBI Search and to the continued development of data discovery services for the "
        "international research community."
    ),
    "Yours faithfully,",
    "",
    "Dr Fatemeh Torabi",
    "Assistant Professor in Health Data Science and Senior Researcher",
    "University of Cambridge and Dementias Platform UK",
]

# Preserve the existing table/contact block, then replace the letter body paragraphs.
for paragraph in doc.paragraphs:
    paragraph.clear()

for idx, paragraph_text in enumerate(body):
    paragraph = doc.paragraphs[idx] if idx < len(doc.paragraphs) else doc.add_paragraph()
    paragraph.clear()
    paragraph.text = paragraph_text

    if idx == 1:
        for run in paragraph.runs:
            run.bold = True
    if idx in {2, 10, 12}:
        for run in paragraph.runs:
            run.bold = True

for paragraph in doc.paragraphs[len(body):]:
    paragraph.clear()

for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(10.5)

doc.save(OUTPUT)
print(OUTPUT)
