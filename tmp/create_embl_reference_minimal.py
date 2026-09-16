from pathlib import Path
import shutil

from docx import Document


ROOT = Path("cvs/recommendation_letters")
SOURCE = ROOT / "Fatemeh_Torabi_General_Academic_Reference_Draft.docx"
OUTPUT = ROOT / "Fatemeh_Torabi_EMBL_EBI_JR3882_Reference_Minimal_Edit.docx"

shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)

replacements = {
    2: "Re: Recommendation for Alieyeh Sarabandi Moghaddam - Software Engineer, Data Discovery (JR3882), EMBL-EBI",
    4: (
        "I am pleased to recommend Ms Alieyeh Sarabandi Moghaddam for the Software Engineer - Data Discovery role "
        "at EMBL-EBI, and for related opportunities in health data science, computational biomedicine and biomedical "
        "informatics. I have worked with Alieyeh as a research and writing supervisor within Dementias Platform UK "
        "(DPUK), and have come to know her as an intellectually curious, conscientious and technically strong researcher."
    ),
    7: (
        "Her broader work at DPUK encompasses reproducible genomic and multi-omic pipelines, cohort-scale polygenic "
        "risk score analysis, dementia subtyping, secure data processing, metadata, provenance and researcher-facing "
        "data discovery. She combines a BSc in Computer Engineering with an MSc in Health Data Science, allowing her "
        "to move confidently between software implementation, data engineering, statistical analysis and biomedical "
        "interpretation. She is especially careful about limitations, generalisability and the governance requirements "
        "of sensitive health data."
    ),
    9: (
        "I recommend Alieyeh with confidence. She would bring strong analytical ability, research maturity and a "
        "collaborative approach to the Software Engineer - Data Discovery role at EMBL-EBI, and I expect her to "
        "contribute thoughtfully to software and data services that support biomedical research."
    ),
}

for idx, text in replacements.items():
    paragraph = doc.paragraphs[idx]
    if not paragraph.runs:
        paragraph.add_run(text)
        continue
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""

doc.save(OUTPUT)
print(OUTPUT)
